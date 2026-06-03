"""
phase2.py — Brand research + presentation package generator.

For each qualified lead in candidates.csv this script:
  1. Pulls brand signals (colors, photos) via Google Places API + Bright Data.
  2. Generates a per-lead sales deck  →  packages/{slug}/deck.pptx  +  deck.pdf
  3. Compiles a master call-script PDF →  packages/master_script.pdf
     (one tabbed section per lead, for use by sales staff)

Usage:
    python phase2.py                         # process all leads in candidates.csv
    python phase2.py --lead "Trevino Drywall"  # single lead by name
    python phase2.py --script-only           # regenerate master script, skip brand/deck
    python phase2.py --deck-only             # regenerate decks, skip script
    python phase2.py --csv other.csv         # use a different input file

Output layout:
    packages/
        master_script.pdf
        trevino-drywall/
            brand.json
            brand_photo.jpg   (if found)
            deck.pptx
            deck.pdf          (exported via PowerPoint COM; skipped if PPT not installed)

Dependencies (add to requirements.txt):
    python-pptx>=0.6.21
    Pillow>=10.0
    reportlab>=4.0
    comtypes>=1.2          # optional — Windows only, used for PPTX→PDF export
"""

from __future__ import annotations

import csv
import io
import json
import os
import re
import shutil
import subprocess
import sys
import textwrap
import time
from collections import Counter
from datetime import datetime
from pathlib import Path
from typing import Any

import requests
from dotenv import load_dotenv

# ---------------------------------------------------------------------------
# Bootstrap — check heavy deps early with a friendly message
# ---------------------------------------------------------------------------

def _require(pkg: str, install_as: str | None = None) -> Any:
    import importlib
    try:
        return importlib.import_module(pkg)
    except ImportError:
        install = install_as or pkg
        sys.exit(
            f"ERROR: '{pkg}' is not installed.\n"
            f"  Run:  pip install {install}\n"
            f"  Then re-run phase2.py."
        )

_pptx_mod   = _require("pptx",      "python-pptx")
_PIL_mod    = _require("PIL",       "Pillow")
_rl_mod     = _require("reportlab", "reportlab")

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt, Emu
from PIL import Image
from reportlab.lib import colors as rl_colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import (
    SimpleDocTemplate, Paragraph, Spacer, HRFlowable,
    PageBreak, Table, TableStyle, KeepTogether
)
from reportlab.platypus.tableofcontents import TableOfContents


# ---------------------------------------------------------------------------
# Config
# ---------------------------------------------------------------------------

load_dotenv()
API_KEY    = os.getenv("GOOGLE_MAPS_API_KEY", "").strip()
_SCRIPT_DIR = Path(__file__).resolve().parent
DEFAULT_CSV = str(_SCRIPT_DIR / "candidates.csv")
PACKAGES_DIR = _SCRIPT_DIR / "packages"

PLACE_PHOTO_URL = "https://places.googleapis.com/v1/{name}/media"
PLACE_DETAILS_URL = "https://places.googleapis.com/v1/places/{place_id}"

SLEEP = 0.4

# Pricing used in decks  — edit as needed
PRICE_STARTER  = "$497"
PRICE_PRO      = "$797"
PRICE_MONTHLY  = "$49/mo"  # hosting + maintenance

# ---------------------------------------------------------------------------
# Category defaults — (primary_rgb, secondary_rgb)
# Used when no brand photo is available.
# ---------------------------------------------------------------------------

_CATEGORY_COLORS: dict[str, tuple[tuple,tuple]] = {
    "house cleaning":      ((30, 144, 255), (255, 255, 255)),
    "carpet cleaning":     ((103, 58, 183), (255, 255, 255)),
    "hvac":                ((13, 71, 161),  (255, 140, 0)),
    "hvac contractor":     ((13, 71, 161),  (255, 140, 0)),
    "plumber":             ((2, 119, 189),  (255, 255, 255)),
    "electrician":         ((245, 127, 23), (33, 33, 33)),
    "painter":             ((255, 87, 34),  (255, 255, 255)),
    "roofer":              ((62, 39, 35),   (255, 193, 7)),
    "handyman":            ((56, 142, 60),  (255, 255, 255)),
    "drywall contractor":  ((96, 125, 139), (255, 255, 255)),
    "pest control":        ((183, 28, 28),  (255, 255, 255)),
    "junk hauling":        ((55, 71, 79),   (255, 152, 0)),
    "junk removal":        ((55, 71, 79),   (255, 152, 0)),
}

def _category_colors(category: str) -> tuple[tuple, tuple]:
    cat = category.lower().strip()
    for k, v in _CATEGORY_COLORS.items():
        if k in cat or cat in k:
            return v
    return ((25, 118, 210), (255, 255, 255))   # default: blue + white


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------

def _slug(name: str) -> str:
    s = name.lower()
    s = re.sub(r"[^a-z0-9]+", "-", s).strip("-")
    return s[:40]


def _hex(rgb: tuple) -> str:
    return "#{:02x}{:02x}{:02x}".format(*rgb)


def _rgb(h: str) -> tuple:
    h = h.lstrip("#")
    return tuple(int(h[i:i+2], 16) for i in (0, 2, 4))


def _pptx_color(rgb: tuple) -> RGBColor:
    return RGBColor(*rgb)


# ---------------------------------------------------------------------------
# Brand research
# ---------------------------------------------------------------------------

def _fetch_place_photos(place_id: str) -> list[str]:
    """Return up to 3 photo resource names from Place Details."""
    if not API_KEY or not place_id:
        return []
    url = PLACE_DETAILS_URL.format(place_id=place_id)
    headers = {
        "X-Goog-Api-Key": API_KEY,
        "X-Goog-FieldMask": "photos",
    }
    try:
        r = requests.get(url, headers=headers, timeout=10)
        if r.status_code != 200:
            return []
        data = r.json()
        photos = data.get("photos") or []
        return [p["name"] for p in photos[:3] if "name" in p]
    except Exception:
        return []


def _download_photo(photo_name: str) -> bytes | None:
    """Download a Places API v1 photo; return raw bytes or None."""
    if not API_KEY:
        return None
    url = f"https://places.googleapis.com/v1/{photo_name}/media"
    params = {"maxHeightPx": 600, "key": API_KEY}
    try:
        r = requests.get(url, params=params, timeout=15, allow_redirects=True)
        if r.status_code == 200 and r.content:
            return r.content
    except Exception:
        pass
    return None


def _extract_colors(img_bytes: bytes) -> tuple[tuple, tuple] | None:
    """Return (primary_rgb, secondary_rgb) from image bytes using Pillow."""
    try:
        img = Image.open(io.BytesIO(img_bytes)).convert("RGB")
        img = img.resize((100, 100), Image.LANCZOS)
        pixels = list(img.getdata())

        # Filter out near-white, near-black, and near-grey (low saturation)
        filtered = []
        for r, g, b in pixels:
            lo, hi = min(r, g, b), max(r, g, b)
            saturation = (hi - lo) / hi if hi > 0 else 0
            if saturation > 0.15 and hi > 40 and lo < 230:
                filtered.append((r // 24 * 24, g // 24 * 24, b // 24 * 24))

        if not filtered:
            return None

        freq = Counter(filtered).most_common(5)
        primary = freq[0][0]
        secondary = freq[1][0] if len(freq) > 1 else (255, 255, 255)
        return primary, secondary
    except Exception:
        return None


def _bd_scrape_social(url: str) -> bytes | None:
    """Try to pull an og:image from a social/website URL via Bright Data CLI."""
    bdata = shutil.which("bdata") or shutil.which("brightdata")
    if not bdata or not url:
        return None
    try:
        proc = subprocess.run(
            [bdata, "scrape", url, "--json"],
            capture_output=True, timeout=60
        )
        if proc.returncode != 0:
            return None
        raw = proc.stdout.decode("utf-8", errors="replace")
        data = json.loads(raw)
        # Look for og:image in the scraped data
        og_image = None
        if isinstance(data, dict):
            og_image = (data.get("og_image") or data.get("image") or
                        data.get("meta", {}).get("og:image"))
        if not og_image:
            return None
        r = requests.get(og_image, timeout=15)
        if r.status_code == 200:
            return r.content
    except Exception:
        pass
    return None


def research_brand(row: dict) -> dict:
    """
    Gather brand intel for one lead. Returns a dict with:
        primary_hex, secondary_hex, photo_path (or None), source
    """
    slug = _slug(row.get("business_name", "unknown"))
    out_dir = PACKAGES_DIR / slug
    out_dir.mkdir(parents=True, exist_ok=True)
    brand_file = out_dir / "brand.json"

    # Return cached result if available
    if brand_file.exists():
        try:
            return json.loads(brand_file.read_text())
        except Exception:
            pass

    primary, secondary = _category_colors(row.get("category", ""))
    photo_path_str = None
    source = "category_default"

    # 1. Try Google Places photos
    place_id = row.get("place_id", "").strip()
    photo_names = _fetch_place_photos(place_id)
    for pname in photo_names:
        img_bytes = _download_photo(pname)
        if not img_bytes:
            continue
        colors = _extract_colors(img_bytes)
        if colors:
            primary, secondary = colors
            # Save the photo
            photo_file = out_dir / "brand_photo.jpg"
            photo_file.write_bytes(img_bytes)
            photo_path_str = str(photo_file)
            source = "google_photo"
            break
        time.sleep(SLEEP)

    # 2. Try scraping their social/website URL via Bright Data (fallback)
    if source == "category_default":
        website_url = row.get("website_url", "").strip()
        if website_url:
            img_bytes = _bd_scrape_social(website_url)
            if img_bytes:
                colors = _extract_colors(img_bytes)
                if colors:
                    primary, secondary = colors
                    photo_file = out_dir / "brand_photo.jpg"
                    photo_file.write_bytes(img_bytes)
                    photo_path_str = str(photo_file)
                    source = "social_scrape"

    brand = {
        "business_name": row.get("business_name", ""),
        "primary_hex":   _hex(primary),
        "secondary_hex": _hex(secondary),
        "photo_path":    photo_path_str,
        "source":        source,
    }
    brand_file.write_text(json.dumps(brand, indent=2))
    print(f"  brand: {primary} / {secondary}  [{source}]")
    return brand


# ---------------------------------------------------------------------------
# PPTX Deck generation
# ---------------------------------------------------------------------------

def _add_slide(prs: Presentation, layout_idx: int = 6):
    layout = prs.slide_layouts[layout_idx]   # Blank layout
    return prs.slides.add_slide(layout)


def _text_box(slide, left, top, width, height,
              text: str, font_size: int, bold=False,
              color_rgb=(255,255,255), align=PP_ALIGN.LEFT, wrap=True):
    from pptx.util import Inches, Pt
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.bold = bold
    run.font.size = Pt(font_size)
    run.font.color.rgb = RGBColor(*color_rgb)
    return txBox


def _fill_bg(slide, rgb: tuple):
    from pptx.util import Inches
    from pptx.dml.color import RGBColor as RC
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = RC(*rgb)


def _accent_bar(slide, primary_rgb, left=0, top=6.8, width=13.33, height=0.4):
    bar = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.RECTANGLE
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    bar.fill.solid()
    bar.fill.fore_color.rgb = RGBColor(*primary_rgb)
    bar.line.fill.background()


def _bullet_box(slide, left, top, width, height,
                items: list[str], font_size: int = 16,
                text_rgb=(255,255,255), bullet="●  "):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = bullet + item
        run.font.size = Pt(font_size)
        run.font.color.rgb = RGBColor(*text_rgb)
    return txBox


def generate_deck(row: dict, brand: dict, out_dir: Path) -> Path:
    """Generate a branded 6-slide PPTX. Returns the path to the saved file."""
    biz       = row.get("business_name", "Unknown Business")
    category  = row.get("category", "local service").title()
    owner     = row.get("possible_owner_name", "").strip()
    stars     = row.get("stars", "")
    reviews   = row.get("review_count", "")
    ws_status = row.get("website_status", "none")
    phone     = row.get("phone", "")
    zip_code  = row.get("zip", "Denver CO")

    primary   = _rgb(brand["primary_hex"])
    dark_bg   = tuple(max(0, c - 40) for c in primary)   # slightly darker variant
    white     = (255, 255, 255)
    near_white = (240, 240, 240)
    dark_text = (30, 30, 30)

    # Website status → human-readable line
    ws_map = {
        "none":       "No website found on Google Business Profile",
        "social_only":"Only a Facebook/Instagram page (no real website)",
        "placeholder":"Auto-generated placeholder site (not a real web presence)",
        "directory":  "Listed on Yelp/Angi — no real website",
        "unreachable":"Website URL exists but is unreachable",
        "broken":     "Website URL is broken (returns error)",
        "blocked":    "Website may exist but blocks visitors",
    }
    ws_note = ws_map.get(ws_status, f"Website status: {ws_status}")

    prs = Presentation()
    prs.slide_width  = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # ------------------------------------------------------------------
    # Slide 1 — Title  (polished split-panel layout)
    # ------------------------------------------------------------------
    s1 = _add_slide(prs)
    _fill_bg(s1, dark_bg)

    # Thin left accent strip
    strip = s1.shapes.add_shape(1, Inches(0), Inches(0), Inches(0.4), Inches(7.5))
    strip.fill.solid(); strip.fill.fore_color.rgb = RGBColor(*primary)
    strip.line.fill.background()

    # Right color panel
    rpanel = s1.shapes.add_shape(
        1, Inches(8.8), Inches(0), Inches(4.53), Inches(7.5)
    )
    rpanel.fill.solid(); rpanel.fill.fore_color.rgb = RGBColor(*primary)
    rpanel.line.fill.background()

    # Brand photo in right panel (if research found one)
    _photo = brand.get("photo_path")
    if _photo:
        try:
            _pp = Path(_photo)
            if _pp.exists():
                s1.shapes.add_picture(
                    str(_pp), Inches(8.9), Inches(0.4), Inches(4.2), Inches(3.4)
                )
        except Exception:
            pass

    # Decorative circle overlay on right panel
    _circ = s1.shapes.add_shape(9, Inches(9.6), Inches(4.0), Inches(2.8), Inches(2.8))
    _lighter = tuple(min(255, c + 35) for c in primary)
    _circ.fill.solid(); _circ.fill.fore_color.rgb = RGBColor(*_lighter)
    _circ.line.fill.background()

    # Business name — left content area
    _text_box(s1, 0.7, 0.9, 8.0, 1.5,
              biz, 40, bold=True, color_rgb=white, align=PP_ALIGN.LEFT)
    _text_box(s1, 0.7, 2.55, 7.6, 0.5,
              category + "  ·  Denver / Aurora, CO", 17,
              color_rgb=near_white, align=PP_ALIGN.LEFT)

    # Thin rule
    _rule = s1.shapes.add_shape(1, Inches(0.7), Inches(3.15), Inches(6.5), Inches(0.04))
    _rule.fill.solid(); _rule.fill.fore_color.rgb = RGBColor(*primary)
    _rule.line.fill.background()

    _text_box(s1, 0.7, 3.3, 7.6, 0.6,
              "Website Proposal", 28, bold=True, color_rgb=white, align=PP_ALIGN.LEFT)
    _text_box(s1, 0.7, 4.05, 7.6, 0.4,
              f"Prepared by Alim's Freelancing  ·  {datetime.today().strftime('%B %Y')}",
              12, color_rgb=near_white, align=PP_ALIGN.LEFT)

    # Rating badge
    if stars:
        _rbadge = s1.shapes.add_shape(
            5, Inches(0.7), Inches(5.4), Inches(3.6), Inches(0.65)
        )
        _rbadge.fill.solid(); _rbadge.fill.fore_color.rgb = RGBColor(*primary)
        _rbadge.line.fill.background()
        _text_box(s1, 0.85, 5.46, 3.4, 0.52,
                  f"⭐  {stars} stars  ·  {reviews} Google reviews",
                  13, bold=True, color_rgb=white)

    # ------------------------------------------------------------------
    # Slide 2 — Where You Are Today
    # ------------------------------------------------------------------
    s2 = _add_slide(prs)
    _fill_bg(s2, (245, 245, 245))
    _accent_bar(s2, primary, top=0, height=0.08)

    _text_box(s2, 0.5, 0.25, 12, 0.7,
              "Where You Are Today", 30, bold=True,
              color_rgb=primary, align=PP_ALIGN.LEFT)

    stats_items = [
        f"⭐  {stars} stars  ·  {reviews} Google reviews — you have social proof.",
        f"📍  {biz}  ·  {zip_code}",
        f"🌐  {ws_note}",
        "👤  " + (f"Owner: {owner}" if owner else "Owner: (to be confirmed)"),
    ]
    _bullet_box(s2, 0.5, 1.2, 12.3, 4.5, stats_items,
                font_size=18, text_rgb=dark_text, bullet="")

    _text_box(s2, 0.5, 6.0, 12, 0.5,
              "You have the reviews. You need the front door.",
              20, bold=True, color_rgb=primary, align=PP_ALIGN.CENTER)
    _accent_bar(s2, primary, top=6.8)

    # ------------------------------------------------------------------
    # Slide 3 — The Gap (opportunity cost)
    # ------------------------------------------------------------------
    s3 = _add_slide(prs)
    _fill_bg(s3, dark_bg)
    _accent_bar(s3, primary, top=0, height=0.08)

    _text_box(s3, 0.5, 0.25, 12, 0.7,
              "What Not Having a Website Costs You", 28, bold=True,
              color_rgb=white, align=PP_ALIGN.LEFT)

    gap_items = [
        "76% of people who search for a local service visit a business within 24 hours.",
        "Businesses with a website get 2× more calls than those without.",
        "AI booking bots capture leads at 2 AM — your voicemail doesn't.",
        f"Competitors in {category.lower()} ARE online. You're invisible on Google Search.",
        "One missed job from a Google search = $200–$800 you didn't capture.",
    ]
    _bullet_box(s3, 0.5, 1.2, 12.3, 5.0, gap_items,
                font_size=17, text_rgb=white)
    _accent_bar(s3, primary, top=6.8)

    # ------------------------------------------------------------------
    # Slide 4 — What We Build
    # ------------------------------------------------------------------
    s4 = _add_slide(prs)
    _fill_bg(s4, (245, 245, 245))
    _accent_bar(s4, primary, top=0, height=0.08)

    _text_box(s4, 0.5, 0.25, 12, 0.7,
              "What We Build For You", 30, bold=True,
              color_rgb=primary, align=PP_ALIGN.LEFT)

    col1 = [
        "Mobile-first design (60%+ of searches are on phones)",
        "AI scheduling bot — customers book 24/7, no calls needed",
        "Google Business Profile sync",
        "Your brand colors, logo, and voice",
        "5-star reviews highlighted front and center",
    ]
    col2 = [
        "Contact form + click-to-call button",
        "Photo gallery of your work",
        "Service area map",
        "Live in 5–7 business days",
        "Annual hosting + maintenance included",
    ]
    _bullet_box(s4, 0.4, 1.2, 6.0, 5.0, col1, font_size=16, text_rgb=dark_text)
    _bullet_box(s4, 6.7, 1.2, 6.0, 5.0, col2, font_size=16, text_rgb=dark_text)
    _accent_bar(s4, primary, top=6.8)

    # ------------------------------------------------------------------
    # Slide 5 — Investment + ROI
    # ------------------------------------------------------------------
    s5 = _add_slide(prs)
    _fill_bg(s5, dark_bg)
    _accent_bar(s5, primary, top=0, height=0.08)

    _text_box(s5, 0.5, 0.25, 12, 0.7,
              "The Investment", 30, bold=True,
              color_rgb=white, align=PP_ALIGN.LEFT)

    tiers = [
        ["Starter", PRICE_STARTER,
         "5-page site + contact form + hosting"],
        ["Pro",     PRICE_PRO,
         "Everything in Starter + AI booking bot + Google Ads-ready"],
        ["Hosting", PRICE_MONTHLY,
         "After Year 1 — hosting, updates, SSL, uptime monitoring"],
    ]
    for i, (name, price, desc) in enumerate(tiers):
        top = 1.4 + i * 1.6
        box = s5.shapes.add_shape(1,
            Inches(0.5), Inches(top), Inches(12.3), Inches(1.3))
        box.fill.solid()
        box.fill.fore_color.rgb = RGBColor(*primary)
        box.line.fill.background()
        _text_box(s5, 0.7, top + 0.08, 3, 0.5,
                  name, 18, bold=True, color_rgb=white)
        _text_box(s5, 3.5, top + 0.08, 2.5, 0.5,
                  price, 22, bold=True, color_rgb=white)
        _text_box(s5, 6.2, top + 0.08, 6.5, 0.5,
                  desc, 15, color_rgb=near_white)

    _text_box(s5, 0.5, 6.0, 12, 0.5,
              "One booked job typically covers 6+ months of cost.",
              18, bold=True, color_rgb=white, align=PP_ALIGN.CENTER)
    _accent_bar(s5, primary, top=6.8)

    # ------------------------------------------------------------------
    # Slide 6 — Next Steps
    # ------------------------------------------------------------------
    s6 = _add_slide(prs)
    _fill_bg(s6, (245, 245, 245))
    _accent_bar(s6, primary, top=0, height=0.08)

    _text_box(s6, 0.5, 0.25, 12, 0.7,
              "Next Steps", 30, bold=True,
              color_rgb=primary, align=PP_ALIGN.LEFT)

    steps = [
        "1.  15-minute Zoom or in-person walkthrough — we show you a mockup using your brand.",
        "2.  You approve the design direction.",
        "3.  We build. You're live in 5–7 days.",
        "4.  AI booking bot goes live. Customers start scheduling online.",
    ]
    _bullet_box(s6, 0.5, 1.3, 12.3, 4.0,
                steps, font_size=19, text_rgb=dark_text, bullet="")

    contact = "alimalhamim@gmail.com  ·  Schedule at: [your calendar link]"
    _text_box(s6, 0.5, 5.8, 12.3, 0.6,
              contact, 15, color_rgb=(100, 100, 100), align=PP_ALIGN.CENTER)
    _accent_bar(s6, primary, top=6.8)

    # Save
    pptx_path = out_dir / "deck.pptx"
    prs.save(str(pptx_path))
    return pptx_path


def export_deck_pdf(pptx_path: Path) -> Path | None:
    """Export PPTX → PDF via PowerPoint COM (Windows only). Returns PDF path or None."""
    pdf_path = pptx_path.with_suffix(".pdf")
    try:
        import comtypes.client  # type: ignore
        ppt = comtypes.client.CreateObject("PowerPoint.Application")
        ppt.Visible = 1
        deck = ppt.Presentations.Open(str(pptx_path.resolve()))
        deck.SaveAs(str(pdf_path.resolve()), 32)  # 32 = ppSaveAsPDF
        deck.Close()
        ppt.Quit()
        return pdf_path
    except Exception:
        return None


# ---------------------------------------------------------------------------
# Master call-script PDF
# ---------------------------------------------------------------------------

def _objection_rows(biz: str, category: str, ws_status: str) -> list[tuple[str,str]]:
    ws_phrase = {
        "social_only":  "Facebook page",
        "placeholder":  "placeholder site",
        "directory":    "Yelp listing",
        "unreachable":  "broken website link",
        "blocked":      "site that blocks visitors",
    }.get(ws_status, "no real website")

    return [
        (
            '"I already have a website / I\'m on Facebook."',
            f"That's great — a lot of our clients started the same way. "
            f"The issue is that a {ws_phrase} doesn't rank on Google when someone searches "
            f"'{category.lower()} near me' — that traffic goes to businesses with their "
            f"own site. We can build something that ties into your existing presence and "
            f"gets you found. Can I show you a quick example?"
        ),
        (
            '"I\'m not interested."',
            "Totally understand — you're busy. Can I ask just one thing: "
            "how are most of your new customers finding you right now? "
            "[Listen. Pivot: 'So most of it is word-of-mouth — that's great, "
            "but it also means if referrals slow down, your pipeline does too. "
            "A site is insurance against that.']"
        ),
        (
            '"How much does it cost?"',
            f"Depends on what you need, but to give you a ballpark — most {category.lower()} "
            f"businesses we work with start at {PRICE_STARTER}. Most find that one job from "
            f"online bookings covers the whole year. That said, I'd need to understand your "
            f"situation before I quote anything — that's really what the 15-minute call is for. "
            f"Can we do [day] at [time]?"
        ),
        (
            '"Send me something / email me."',
            "Absolutely — what's the best email? [get it] And just so I send you "
            "something relevant: what's the main thing you're trying to grow — new "
            "residential clients, commercial, or both? [get their answer + intel] "
            "I'll send that today. Can we pencil in a quick 15-minute call for "
            "[day] at [time] to walk through it together?"
        ),
        (
            '"I\'ll think about it / call me back later."',
            "Of course. Can I ask what you'd be thinking about specifically? "
            "[Listen and address.] One thing I'll mention — we're booking builds "
            "for [next month] and slots fill up. But I understand you want to think. "
            "When's the best time to follow up — end of this week?"
        ),
        (
            '"I do my own marketing / a family member helps."',
            "That's awesome — having someone in-house is huge. What we do is "
            "build the foundation: the site, the booking system, the Google SEO. "
            "Your person can keep running the social side and it all works together. "
            "It's less work for them, not more. Worth 15 minutes to see the setup?"
        ),
    ]


def _build_script_story(rows: list[dict], brands: dict[str, dict],
                        styles: dict) -> list:
    """Build a list of ReportLab flowables for the master script."""
    story = []

    # Cover page
    story.append(Spacer(1, 1.5 * inch))
    story.append(Paragraph(
        "<b>Sales Call Script</b>",
        styles["CoverTitle"]
    ))
    story.append(Spacer(1, 0.2 * inch))
    story.append(Paragraph(
        f"Alim's Freelancing  ·  {len(rows)} Qualified Leads  ·  "
        f"{datetime.today().strftime('%B %Y')}",
        styles["CoverSub"]
    ))
    story.append(Spacer(1, 0.5 * inch))
    story.append(Paragraph(
        "<b>How to use this script</b>",
        styles["SectionHeader"]
    ))
    usage = (
        "Each lead has its own section below. Read the opener naturally — "
        "don't recite it word-for-word. The rebuttal tree covers the six "
        "most common objections; trust it but adapt to the conversation. "
        "The goal of every call is ONE ask: a 15-minute walkthrough, "
        "Zoom or in-person."
    )
    story.append(Paragraph(usage, styles["BodyText"]))
    story.append(PageBreak())

    for row in rows:
        biz       = row.get("business_name", "Unknown")
        owner     = row.get("possible_owner_name", "").strip()
        category  = row.get("category", "local service").title()
        phone     = row.get("phone", "")
        stars     = row.get("stars", "")
        reviews   = row.get("review_count", "")
        ws_status = row.get("website_status", "none")
        address   = row.get("address", "")
        slug      = _slug(biz)
        brand     = brands.get(slug, {})
        primary_hex = brand.get("primary_hex", "#1976d2")

        # Lead header bar
        story.append(KeepTogether([
            Paragraph(f"<b>{biz}</b>", styles["LeadTitle"]),
            Paragraph(
                f"📞 {phone}   &nbsp;&nbsp;  ⭐ {stars} ({reviews} reviews)   "
                f"&nbsp;&nbsp;  🏷 {category}   &nbsp;&nbsp;  📍 {address}",
                styles["LeadMeta"]
            ),
            Paragraph(
                f"Owner: <b>{owner or '(confirm on call)'}</b>   &nbsp;&nbsp;  "
                f"Website: <b>{ws_status}</b>",
                styles["LeadMeta"]
            ),
            HRFlowable(width="100%", thickness=2, color=rl_colors.HexColor(primary_hex)),
            Spacer(1, 0.15 * inch),
        ]))

        # Opener
        story.append(Paragraph("<b>OPENER</b>", styles["ScriptLabel"]))
        if owner:
            opener_name = f"<b>{owner}</b>"
        else:
            opener_name = "the owner"

        opener = (
            f"\"Hi, is this {opener_name}? / Hi, can I speak with the owner? "
            f"This is [your name] with Alim's Freelancing. I'm calling because "
            f"I came across {biz} on Google — {stars} stars with {reviews} reviews, "
            f"which is impressive — but I couldn't find a proper website for you. "
            f"We work with {category.lower()} businesses in the Denver area to get "
            f"them online and set up with AI scheduling so new customers can book "
            f"without calling. Do you have 2 minutes?\""
        )
        story.append(Paragraph(opener, styles["ScriptText"]))
        story.append(Spacer(1, 0.15 * inch))

        # Pitch
        story.append(Paragraph("<b>PITCH  (after 'yes, go ahead')</b>", styles["ScriptLabel"]))
        pitch = (
            f"\"Great. What we do is build professional websites — custom to "
            f"your brand, your colors, your work — with an AI booking bot built in. "
            f"So when someone searches '{category.lower()} near me' in Denver, "
            f"they find you, and they can schedule right there from their phone. "
            f"No voicemail tag. Most of our clients in {category.lower()} see their "
            f"first online booking within the first week. We set everything up in "
            f"5 to 7 days and you don't have to do anything technical. "
            f"I'd love to show you what it would look like specifically for {biz}. "
            f"Can we do a 15-minute Zoom this week — or I can swing by?\""
        )
        story.append(Paragraph(pitch, styles["ScriptText"]))
        story.append(Spacer(1, 0.2 * inch))

        # Objections
        story.append(Paragraph("<b>OBJECTION HANDLERS</b>", styles["ScriptLabel"]))
        story.append(Spacer(1, 0.08 * inch))

        objections = _objection_rows(biz, category, ws_status)
        tbl_data = [["Objection", "Response"]]
        for obj, resp in objections:
            tbl_data.append([
                Paragraph(f"<i>{obj}</i>", styles["ObjQ"]),
                Paragraph(f"\"{resp}\"", styles["ObjA"]),
            ])

        tbl = Table(tbl_data, colWidths=[2.2 * inch, 5.3 * inch])
        tbl.setStyle(TableStyle([
            ("BACKGROUND",   (0, 0), (-1, 0),  rl_colors.HexColor(primary_hex)),
            ("TEXTCOLOR",    (0, 0), (-1, 0),  rl_colors.white),
            ("FONTNAME",     (0, 0), (-1, 0),  "Helvetica-Bold"),
            ("FONTSIZE",     (0, 0), (-1, 0),  9),
            ("ROWBACKGROUNDS", (0, 1), (-1, -1),
             [rl_colors.HexColor("#f5f5f5"), rl_colors.white]),
            ("VALIGN",       (0, 0), (-1, -1),  "TOP"),
            ("GRID",         (0, 0), (-1, -1),  0.5, rl_colors.HexColor("#cccccc")),
            ("LEFTPADDING",  (0, 0), (-1, -1),  6),
            ("RIGHTPADDING", (0, 0), (-1, -1),  6),
            ("TOPPADDING",   (0, 0), (-1, -1),  5),
            ("BOTTOMPADDING",(0, 0), (-1, -1),  5),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 0.2 * inch))

        # Close
        story.append(Paragraph("<b>THE CLOSE</b>", styles["ScriptLabel"]))
        close = (
            "\"Great — I'll put you down for [day/time]. "
            "You'll get a Zoom link. All I ask is that you show up "
            "and tell me honestly whether it makes sense for you. Fair?\""
        )
        story.append(Paragraph(close, styles["ScriptText"]))
        story.append(PageBreak())

    return story


def generate_master_script(rows: list[dict], brands: dict[str, dict]) -> Path:
    """Generate the master PDF call script for all leads."""
    PACKAGES_DIR.mkdir(parents=True, exist_ok=True)
    pdf_path = PACKAGES_DIR / "master_script.pdf"

    # ReportLab styles
    base = getSampleStyleSheet()

    def S(name, parent="Normal", **kw):
        return ParagraphStyle(name, parent=base[parent], **kw)

    styles = {
        "CoverTitle": S("CoverTitle", "Heading1",
                        fontSize=28, spaceAfter=10, alignment=1),
        "CoverSub":   S("CoverSub", "Normal",
                        fontSize=13, textColor=rl_colors.grey, alignment=1,
                        spaceAfter=18),
        "SectionHeader": S("SectionHeader", "Heading2",
                           fontSize=13, spaceBefore=10, spaceAfter=6),
        "BodyText":   S("BodyText", "Normal",
                        fontSize=10, leading=14, spaceAfter=8),
        "LeadTitle":  S("LeadTitle", "Heading1",
                        fontSize=18, spaceBefore=4, spaceAfter=2,
                        textColor=rl_colors.HexColor("#1a1a2e")),
        "LeadMeta":   S("LeadMeta", "Normal",
                        fontSize=9, textColor=rl_colors.grey, spaceAfter=3),
        "ScriptLabel": S("ScriptLabel", "Normal",
                         fontSize=9, textColor=rl_colors.grey,
                         spaceBefore=6, spaceAfter=3,
                         fontName="Helvetica-Bold"),
        "ScriptText": S("ScriptText", "Normal",
                        fontSize=10, leading=15, spaceAfter=6,
                        backColor=rl_colors.HexColor("#f0f4ff"),
                        borderPadding=6),
        "ObjQ": S("ObjQ", "Normal", fontSize=9, leading=12),
        "ObjA": S("ObjA", "Normal", fontSize=9, leading=13),
    }

    doc = SimpleDocTemplate(
        str(pdf_path),
        pagesize=letter,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
    )

    story = _build_script_story(rows, brands, styles)
    doc.build(story)
    return pdf_path


# ---------------------------------------------------------------------------
# Load leads
# ---------------------------------------------------------------------------

def load_leads(csv_path: str) -> list[dict]:
    """Load all rows from candidates.csv. Returns all rows (Phase 2 runs on all)."""
    p = Path(csv_path)
    if not p.exists():
        sys.exit(f"ERROR: {csv_path} not found. Run lead_finder.py first.")
    with p.open("r", newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


# ---------------------------------------------------------------------------
# Entry point
# ---------------------------------------------------------------------------

def main():
    import argparse

    parser = argparse.ArgumentParser(
        description=(
            "phase2.py — Brand research + presentation package generator.\n"
            "\n"
            "Reads candidates.csv, then for each lead:\n"
            "  1. Pulls brand colors from Google Business photos (or BD scrape).\n"
            "  2. Generates a branded sales deck  →  packages/{slug}/deck.pptx + .pdf\n"
            "  3. Compiles master call script       →  packages/master_script.pdf\n"
        ),
        formatter_class=argparse.RawDescriptionHelpFormatter,
    )
    parser.add_argument("--csv",         default=DEFAULT_CSV, metavar="PATH",
                        help="Path to candidates.csv (default: next to script)")
    parser.add_argument("--lead",        metavar="NAME",
                        help="Process a single lead by business name")
    parser.add_argument("--script-only", action="store_true",
                        help="Skip brand research and deck gen; regenerate script only")
    parser.add_argument("--deck-only",   action="store_true",
                        help="Skip script; regenerate decks only")
    parser.add_argument("--no-pdf",      action="store_true",
                        help="Skip PPTX→PDF export (useful if PowerPoint not installed)")
    args = parser.parse_args()

    rows = load_leads(args.csv)

    if args.lead:
        needle = args.lead.lower().strip()
        rows = [r for r in rows if needle in r.get("business_name", "").lower()]
        if not rows:
            sys.exit(f"No lead matching '{args.lead}' found in {args.csv}")

    print(f"\nPhase 2  ·  {len(rows)} leads  ·  output → {PACKAGES_DIR}\n")
    PACKAGES_DIR.mkdir(parents=True, exist_ok=True)

    brands: dict[str, dict] = {}

    # Step 1 — brand research
    if not args.script_only:
        print("─" * 50)
        print("STEP 1 — Brand research")
        print("─" * 50)
        for i, row in enumerate(rows, 1):
            biz = row.get("business_name", "?")
            slug = _slug(biz)
            print(f"[{i}/{len(rows)}] {biz}")
            brand = research_brand(row)
            brands[slug] = brand
            time.sleep(SLEEP)
    else:
        # Load any existing brand data from cache
        for row in rows:
            slug = _slug(row.get("business_name", ""))
            brand_file = PACKAGES_DIR / slug / "brand.json"
            if brand_file.exists():
                try:
                    brands[slug] = json.loads(brand_file.read_text())
                except Exception:
                    brands[slug] = {}

    # Step 2 — deck generation
    if not args.script_only:
        print()
        print("─" * 50)
        print("STEP 2 — Deck generation")
        print("─" * 50)
        pdf_missing = []
        for i, row in enumerate(rows, 1):
            biz  = row.get("business_name", "?")
            slug = _slug(biz)
            out  = PACKAGES_DIR / slug
            out.mkdir(parents=True, exist_ok=True)
            brand = brands.get(slug, {})
            if not brand:
                brand = {
                    "primary_hex":   _hex(_category_colors(row.get("category",""))[0]),
                    "secondary_hex": _hex(_category_colors(row.get("category",""))[1]),
                }
            print(f"[{i}/{len(rows)}] {biz}  →  {out}/deck.pptx")
            pptx_path = generate_deck(row, brand, out)

            if not args.no_pdf:
                pdf_path = export_deck_pdf(pptx_path)
                if pdf_path:
                    print(f"        deck.pdf  ✓")
                else:
                    pdf_missing.append(biz)
                    print(f"        deck.pdf  — PowerPoint not found; open deck.pptx and export manually")

        if pdf_missing:
            print(f"\nNote: PDF export requires Microsoft PowerPoint installed.")
            print(f"Open each deck.pptx → File → Export → Create PDF/XPS to export manually.")

    # Step 3 — master script
    if not args.deck_only:
        print()
        print("─" * 50)
        print("STEP 3 — Master call script")
        print("─" * 50)

        # Rebuild brands dict if we skipped research
        if args.script_only:
            for row in rows:
                slug = _slug(row.get("business_name", ""))
                if slug not in brands:
                    brands[slug] = {
                        "primary_hex": _hex(_category_colors(row.get("category",""))[0]),
                    }

        script_path = generate_master_script(rows, brands)
        print(f"master_script.pdf  →  {script_path}")

    print()
    print("=" * 50)
    print("Phase 2 complete.")
    print(f"Output: {PACKAGES_DIR}")
    print("=" * 50)


if __name__ == "__main__":
    main()
